"""Append new rules to a copy of the user's ``rule_catalogue.xlsx`` template.

* The template is copied, never edited in place; existing rows are untouched.
* New rows continue the template's own ``RULE_NNN`` numbering; the returned map (hash id -> final id) keeps
  ``executable_rules.jsonl`` joined to the workbook.
* Every cell is neutralised against spreadsheet formula injection: rule text is model output derived from
  untrusted documents, and openpyxl would otherwise store a string starting with ``=`` as a live formula.
* Optional AI/copyright footer (print footer + on-screen note) - configured in ``configs/base.yaml``.
"""
from __future__ import annotations

import os
import re
import tempfile
from copy import copy
from pathlib import Path
from typing import Any, Sequence

from ..domain.schema import EXCEL_COLUMNS, Rule
from ..errors import OutputError

_RULE_ID = re.compile(r"^RULE_(\d+)$")
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_TEXT_COLUMNS = {"rule_logic_plain_english", "upcoding_mechanism", "source_reference", "rule_name"}


def sanitize_cell(value: Any) -> Any:
    """OWASP-style CSV/Excel injection defence: neutralise values a spreadsheet could interpret as a formula."""
    if isinstance(value, str) and value.startswith(_DANGEROUS_PREFIXES):
        return "'" + value
    return value


def _check_header(ws) -> None:
    header = [str(c.value or "").strip().lower() for c in ws[1][: len(EXCEL_COLUMNS)]]
    expected = [c.lower() for c in EXCEL_COLUMNS]
    if header != expected:
        raise OutputError("template header row does not match the expected 11 columns",
                          details=[f"expected {expected}", f"found    {header}"])


def _last_rule_row(ws) -> tuple[int, int]:
    last, max_n = 1, 0
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=1).value
        m = _RULE_ID.match(str(value)) if value is not None else None
        if m:
            last = row
            max_n = max(max_n, int(m.group(1)))
    return last, max_n


def _copy_style(src, dst) -> None:
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.number_format, dst.protection = copy(src.alignment), src.number_format, copy(src.protection)


def append_rules_to_template(template_path: Path | str, out_path: Path | str, rules: Sequence[Rule], *,
                             sheet_name: str = "rule_catalogue", footer_lines: Sequence[str] | None = None
                             ) -> dict[str, str]:
    """Write ``out_path`` = template + ``rules``; return ``{rule.rule_id: final RULE_NNN}``."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font

    template = Path(template_path)
    if not template.is_file():
        raise OutputError(f"template workbook not found: {template.name}")
    wb = load_workbook(template)
    if sheet_name not in wb.sheetnames:
        raise OutputError(f"sheet {sheet_name!r} not found in template (has {wb.sheetnames})")
    ws = wb[sheet_name]
    _check_header(ws)

    footer = [str(line) for line in (footer_lines or ()) if str(line).strip()]
    last_row, max_n = _last_rule_row(ws)
    for row in range(ws.max_row, last_row, -1):  # drop a footer note left by a previous generation
        if str(ws.cell(row=row, column=1).value or "").strip() in footer:
            ws.delete_rows(row)

    style_row = ws[last_row] if last_row >= 2 else None
    id_map: dict[str, str] = {}
    for offset, rule in enumerate(rules, start=1):
        final_id = f"RULE_{max_n + offset:03d}"
        id_map[rule.rule_id] = final_id
        values = rule.as_row()
        values["rule_id"] = final_id
        target_row = last_row + offset
        for col_idx, column in enumerate(EXCEL_COLUMNS, start=1):
            cell = ws.cell(row=target_row, column=col_idx)
            cell.value = sanitize_cell(values[column])
            cell.data_type = "s"
            if style_row is not None:
                _copy_style(style_row[col_idx - 1], cell)
            else:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(wrap_text=column in _TEXT_COLUMNS, vertical="top")
    end_row = last_row + len(rules)

    if footer:
        gray = Font(name="Arial", size=8, italic=True, color="808080")
        for i, line in enumerate(footer):
            cell = ws.cell(row=end_row + 3 + i, column=1, value=sanitize_cell(line))
            cell.font, cell.alignment = gray, Alignment(wrap_text=False)
        ws.oddFooter.center.text = "\n".join(line.replace("&", "&&") for line in footer)
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.color = "808080"

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".catalogue-", suffix=".xlsx", dir=str(out.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        check = load_workbook(tmp)[sheet_name]
        if _last_rule_row(check)[0] != end_row:
            raise OutputError("post-write verification failed: unexpected row count")
        os.replace(tmp, out)
    except BaseException:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise
    return id_map
