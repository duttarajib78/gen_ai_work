import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from rule_extractor.domain.schema import EXCEL_COLUMNS, validate_rule_dict
from rule_extractor.errors import OutputError
from rule_extractor.output.excel_writer import append_rules_to_template, sanitize_cell

FOOTER = ["Footer line one.", "Footer line two."]


def make_template(path, rows=3, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "rule_catalogue"
    ws.append(header or EXCEL_COLUMNS)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True)
    ws.freeze_panes = "A2"
    for i in range(1, rows + 1):
        ws.append([f"RULE_{i:03d}", f"name {i}", "logic", "A00", "N/A", "N/A", "mech", "remove_code", "p.1", "high", "verified"])
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def rule(name, logic="logic text", **kw):
    base = dict(rule_name=name, rule_logic_plain_english=logic, upcoding_mechanism="mech",
                recommended_action="remove_code", source_reference="icd10am_tabular printed p.27")
    base.update(kw)
    return validate_rule_dict(base)


class ExcelWriterTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.template = self.dir / "template.xlsx"
        make_template(self.template)

    def tearDown(self):
        self.tmp.cleanup()

    def test_appends_continuing_numbering_and_preserves_existing_rows(self):
        out = self.dir / "out.xlsx"
        rules = [rule("a", "one"), rule("b", "two")]
        id_map = append_rules_to_template(self.template, out, rules)
        ws = load_workbook(out)["rule_catalogue"]
        self.assertEqual([ws.cell(row=r, column=1).value for r in range(2, 7)],
                         ["RULE_001", "RULE_002", "RULE_003", "RULE_004", "RULE_005"])
        self.assertEqual(ws["B2"].value, "name 1")
        self.assertEqual(ws["K2"].value, "verified")             # human sign-off in the template is untouched
        self.assertEqual(ws["K5"].value, "needs_review")         # agent rows never 'verified'
        self.assertEqual(id_map[rules[0].rule_id], "RULE_004")
        self.assertEqual(ws["B5"].font.name, "Arial")             # style copied from the last template row
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(load_workbook(self.template)["rule_catalogue"].max_row, 4)   # template never modified

    def test_formula_injection_is_neutralised(self):
        out = self.dir / "out.xlsx"
        evil = rule("=HYPERLINK(\"http://evil\",\"x\")", "+cmd|' /C calc'!A0", source_reference="@SUM(1+1) p.1")
        append_rules_to_template(self.template, out, [evil])
        ws = load_workbook(out)["rule_catalogue"]
        for cell in ("B5", "C5", "I5"):
            self.assertNotEqual(ws[cell].data_type, "f", cell)
            self.assertTrue(str(ws[cell].value).startswith("'"), cell)
        self.assertEqual(sanitize_cell("plain"), "plain")
        self.assertEqual(sanitize_cell("-1"), "'-1")

    def test_footer_print_and_on_screen_without_duplicates_on_rerun(self):
        out1, out2 = self.dir / "o1.xlsx", self.dir / "o2.xlsx"
        append_rules_to_template(self.template, out1, [rule("a")], footer_lines=FOOTER)
        ws = load_workbook(out1)["rule_catalogue"]
        self.assertEqual(ws["A8"].value, FOOTER[0])
        self.assertEqual(ws["A9"].value, FOOTER[1])
        self.assertIn(FOOTER[0], ws.oddFooter.center.text)
        append_rules_to_template(out1, out2, [rule("b", "different logic")], footer_lines=FOOTER)
        ws2 = load_workbook(out2)["rule_catalogue"]
        values = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
        self.assertEqual(values.count(FOOTER[0]), 1)
        self.assertEqual(ws2["A6"].value, "RULE_005")

    def test_header_mismatch_and_missing_sheet_fail_closed(self):
        bad = self.dir / "bad.xlsx"
        make_template(bad, header=["id"] + EXCEL_COLUMNS[1:])
        with self.assertRaises(OutputError):
            append_rules_to_template(bad, self.dir / "x.xlsx", [rule("a")])
        with self.assertRaises(OutputError):
            append_rules_to_template(self.template, self.dir / "x.xlsx", [rule("a")], sheet_name="nope")
        with self.assertRaises(OutputError):
            append_rules_to_template(self.dir / "missing.xlsx", self.dir / "x.xlsx", [])

    def test_empty_template_gets_default_style_and_first_id(self):
        empty = self.dir / "empty.xlsx"
        make_template(empty, rows=0)
        out = self.dir / "o.xlsx"
        append_rules_to_template(empty, out, [rule("a")])
        ws = load_workbook(out)["rule_catalogue"]
        self.assertEqual(ws["A2"].value, "RULE_001")
        self.assertEqual(ws["B2"].font.name, "Arial")


if __name__ == "__main__":
    unittest.main()
