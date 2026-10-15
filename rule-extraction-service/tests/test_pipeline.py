import asyncio
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from rule_extractor.agents.types import AgentRunStats
from rule_extractor.domain.api import JobOptions
from rule_extractor.domain.jobs import JobRecord, JobStatus, UnitRecord, UnitStatus
from rule_extractor.domain.schema import validate_rule_dict
from rule_extractor.errors import AgentError, ConfigError, PolicyError, RequestValidationError, UnavailableError
from rule_extractor.ingestion.corpus import CorpusIndex
from rule_extractor.output.jsonl import iter_jsonl
from rule_extractor.pipeline.dedupe import dedupe_rules
from rule_extractor.pipeline.finalize import Finalizer
from rule_extractor.pipeline.planner import build_plan
from rule_extractor.pipeline.processor import UnitProcessor, unit_rules_key
from rule_extractor.storage.memory import InMemoryArtifactStore

from .fakes import FakeRunner, default_rule_args, emit_default
from .helpers import make_env


def run(coro):
    return asyncio.run(coro)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.settings = make_env(self.tmp.name)
        self.corpus = CorpusIndex(self.settings.paths.corpus_dir)


class PlannerTests(Base):
    def plan(self, **opts):
        return build_plan(JobOptions.from_dict(opts), self.settings, self.corpus)

    def test_default_plan_matches_corpus(self):
        plan = self.plan()
        self.assertEqual(plan.summary["acs_topics"], 6)
        self.assertEqual(plan.summary["sweep_windows"], {"icd10am_tabular": 1, "achi_tabular": 1, "icd10am_alphabetic": 1, "achi_alphabetic": 1})
        self.assertEqual(len(plan.units), 10)
        self.assertEqual([u.order for u in plan.units], list(range(10)))
        self.assertEqual(len({u.unit_id for u in plan.units}), 10)

    def test_filters_and_limits(self):
        plan = self.plan(topics=["0110"], include_tabular_sweep=False)
        self.assertEqual([(u.unit_id, u.spec["start_page"], u.spec["end_page"]) for u in plan.units], [("acs-0110", 6, 7)])
        self.assertEqual(len(self.plan(limit_topics=2, include_tabular_sweep=False).units), 2)
        self.assertEqual(len(self.plan(include_acs_topics=False, sweep_documents=["achi_tabular"]).units), 1)
        self.assertEqual(self.plan(include_acs_topics=False, limit_sweep_windows=1).summary["sweep_windows_total"], 4)

    def test_invalid_plans(self):
        with self.assertRaises(RequestValidationError):
            self.plan(topics=["9999"])
        with self.assertRaises(RequestValidationError):
            self.plan(include_acs_topics=False, include_tabular_sweep=False)
        with self.assertRaises(PolicyError):
            self.plan(include_web_research=True, web_research_topics=["Coding rules"])
        with self.assertRaises(UnavailableError):
            build_plan(JobOptions.from_dict({}), self.settings, None)

    def test_web_research_when_enabled(self):
        enabled = make_env(Path(self.tmp.name) / "w", RULE_EXTRACTOR_ENABLE_WEB_RESEARCH="true")
        with self.assertRaises(RequestValidationError):                       # no default topics configured
            build_plan(JobOptions.from_dict({"include_web_research": True}), enabled, self.corpus)
        plan = build_plan(JobOptions.from_dict({"include_web_research": True, "web_research_topics": ["Coding rules"],
                                                "include_acs_topics": False, "include_tabular_sweep": False}), enabled, self.corpus)
        self.assertEqual(plan.units[0].kind, "web_research")

    def test_unit_cap_and_contract_plan(self):
        capped = make_env(Path(self.tmp.name) / "c", RULE_EXTRACTOR_MAX_UNITS_PER_JOB="3")
        with self.assertRaises(RequestValidationError):
            build_plan(JobOptions.from_dict({}), capped, self.corpus)
        plan = build_plan(JobOptions.from_dict({"kind": "contract_check"}), self.settings, None)
        self.assertEqual([(u.unit_id, u.kind) for u in plan.units], [("contract-check", "contract_check")])


class DedupeTests(unittest.TestCase):
    def test_first_wins_but_executable_replaces_narrative(self):
        def r(name, logic):
            return validate_rule_dict(dict(rule_name=name, rule_logic_plain_english=logic, upcoding_mechanism="m",
                                           recommended_action="remove_code", source_reference="p.1"))
        from rule_extractor.engine.rules_store import build_executable
        a, b, c = r("a", "Same   LOGIC"), r("b", "same logic"), r("c", "other")
        exe, _ = build_executable(b.rule_id, "remove_code", {"has_diagnosis": "A00"}, {"action": "remove_code", "target": "diagnosis", "match": "A00"})
        kept, dropped = dedupe_rules([(a, None), (c, None), (b, exe)])
        self.assertEqual(dropped, 1)
        self.assertEqual([k[0].rule_name for k in kept], ["b", "c"])      # executable duplicate took the first slot
        kept, dropped = dedupe_rules([(a, exe), (b, None)])
        self.assertEqual(kept[0][0].rule_name, "a")


def job_unit(settings, corpus, unit_id="acs-0110"):
    plan = build_plan(JobOptions.from_dict({}), settings, corpus)
    spec = next(u for u in plan.units if u.unit_id == unit_id)
    job = JobRecord(job_id="job-aaaaaaaaaaaa")
    return job, UnitRecord(job.job_id, spec.unit_id, spec.kind, spec.order, spec.spec)


class ProcessorTests(Base):
    def processor(self, runner, sleeps=None):
        self.artifacts = InMemoryArtifactStore()

        async def fake_sleep(delay):
            (sleeps if sleeps is not None else []).append(delay)

        return UnitProcessor(self.settings, self.corpus, runner, self.artifacts, sleep=fake_sleep, rng=lambda: 0.5)

    def test_acs_unit_builds_corpus_task_and_persists_rules(self):
        runner = FakeRunner()
        job, unit = job_unit(self.settings, self.corpus)
        out = run(self.processor(runner).process(job, unit))
        task = runner.tasks[0]
        self.assertEqual(task.profile.builtin_tools, ("Read", "Grep", "Glob"))
        self.assertEqual(task.profile.file_root, Path(self.corpus.root).resolve())
        self.assertEqual(task.model, "au.anthropic.claude-sonnet-5")
        self.assertTrue(task.untrusted_text_in_prompt)
        self.assertIn("Example 4", task.user_prompt)
        self.assertEqual((out.rules_count, out.cost_usd, out.turns, out.rules_key), (1, 0.05, 3, unit_rules_key(job.job_id, unit.unit_id)))
        record = next(iter_jsonl(self.artifacts.get_bytes(out.rules_key)))
        self.assertEqual(record["rule"]["job_id"], job.job_id)
        self.assertEqual(record["rule"]["topic"], "ACS 0110 Bacterial, viral and other infectious agents")
        self.assertEqual(record["rule"]["confidence"], "needs_review")
        self.assertIsNotNone(record["executable"])

    def test_sweep_units_have_no_tools_and_model_follows_document_kind(self):
        runner = FakeRunner(lambda t, n: AgentRunStats(turns=1))
        for unit_id, model in (("sweep-icd10am_tabular-w0000", "au.anthropic.claude-sonnet-5"),
                               ("sweep-achi_alphabetic-w0000", "au.anthropic.claude-haiku-4-5-20251001-v1:0")):
            job, unit = job_unit(self.settings, self.corpus, unit_id)
            out = run(self.processor(runner).process(job, unit))
            self.assertEqual(runner.tasks[-1].profile.builtin_tools, ())
            self.assertEqual(runner.tasks[-1].model, model)
            self.assertEqual(out.rules_count, 0)
            self.assertEqual(self.artifacts.get_bytes(out.rules_key), b"")          # zero rules is a valid, recorded outcome

    def test_transient_failures_retry_with_backoff_and_escalating_turns(self):
        def script(task, n):
            if n <= 2:
                raise AgentError("hit turns", transient=True, reason="max_turns", stats=AgentRunStats(cost_usd=0.5))
            return emit_default(task)

        sleeps, runner = [], FakeRunner(script)
        job, unit = job_unit(self.settings, self.corpus)
        out = run(self.processor(runner, sleeps).process(job, unit))
        self.assertEqual([t.max_turns for t in runner.tasks], [40, 61, 92])
        self.assertEqual(sleeps, [0.0, 0.0])                                        # base delay is 0 in this env
        self.assertAlmostEqual(out.cost_usd, 1.05)                                  # failed attempts are still paid for
        self.assertGreater(runner.tasks[2].timeout_seconds, runner.tasks[0].timeout_seconds)

    def test_backoff_is_exponential_with_jitter_and_capped(self):
        settings = make_env(Path(self.tmp.name) / "b", RULE_EXTRACTOR_RETRY_BASE_DELAY_SECONDS="5", RULE_EXTRACTOR_MAX_RETRIES="3")
        runner = FakeRunner(lambda t, n: (_ for _ in ()).throw(AgentError("throttled", transient=True)) if n < 4 else emit_default(t))
        sleeps = []

        async def fake_sleep(d):
            sleeps.append(d)

        proc = UnitProcessor(settings, self.corpus, runner, InMemoryArtifactStore(), sleep=fake_sleep, rng=lambda: 0.5)
        job, unit = job_unit(settings, self.corpus)
        run(proc.process(job, unit))
        self.assertEqual(sleeps, [5.0, 10.0, 20.0])

    def test_non_transient_failure_raises_immediately_with_cost(self):
        runner = FakeRunner(lambda t, n: (_ for _ in ()).throw(AgentError("bad request", transient=False, stats=AgentRunStats(cost_usd=0.3))))
        job, unit = job_unit(self.settings, self.corpus)
        with self.assertRaises(AgentError) as cm:
            run(self.processor(runner).process(job, unit))
        self.assertEqual(len(runner.tasks), 1)
        self.assertEqual(cm.exception.cost_usd, 0.3)

    def test_retries_exhaust(self):
        runner = FakeRunner(lambda t, n: (_ for _ in ()).throw(AgentError("throttled", transient=True)))
        job, unit = job_unit(self.settings, self.corpus)
        with self.assertRaises(AgentError):
            run(self.processor(runner).process(job, unit))
        self.assertEqual(len(runner.tasks), self.settings.agent.max_retries + 1)

    def test_partial_rules_from_failed_attempt_are_discarded(self):
        def script(task, n):
            emit_default(task, count=2, logic=f"attempt {n}")       # emit then fail
            if n == 1:
                raise AgentError("crashed late", transient=True)
            return AgentRunStats(turns=2, cost_usd=0.01)

        runner = FakeRunner(script)
        job, unit = job_unit(self.settings, self.corpus)
        out = run(self.processor(runner).process(job, unit))
        self.assertEqual(out.rules_count, 2)
        stored = self.artifacts.get_bytes(out.rules_key).decode()
        self.assertIn("attempt 2", stored)
        self.assertNotIn("attempt 1", stored)                         # nothing from the failed attempt leaks in

    def test_web_unit_requires_enabled_config(self):
        job = JobRecord(job_id="job-aaaaaaaaaaaa")
        unit = UnitRecord(job.job_id, "web-x", "web_research", 0, {"topic": "Coding rules"})
        with self.assertRaises(ConfigError):
            run(self.processor(FakeRunner()).process(job, unit))
        enabled = make_env(Path(self.tmp.name) / "w", RULE_EXTRACTOR_ENABLE_WEB_RESEARCH="true")
        runner = FakeRunner(lambda t, n: AgentRunStats(turns=1))
        proc = UnitProcessor(enabled, self.corpus, runner, InMemoryArtifactStore())
        run(proc.process(job, unit))
        self.assertEqual(runner.tasks[0].profile.builtin_tools, ("WebFetch",))
        self.assertEqual(runner.tasks[0].profile.web_allowed_hosts, ("www.ihacpa.gov.au", "ihacpa.gov.au"))
        self.assertFalse(runner.tasks[0].untrusted_text_in_prompt)

    def test_contract_check_unit(self):
        def good(task, n):
            names = {t.name for t in task.tools}
            if "record_simple" in names:
                task.tools[0].handler({"message": "hello world"})
            elif "record_nested" in names:
                task.tools[0].handler({"message": "hello world", "detail": {"a": 1}})
            else:
                return AgentRunStats(turns=1, denied_tools=["Read"], tool_calls={"Read": 1})
            return AgentRunStats(turns=1, cost_usd=0.01)

        job = JobRecord(job_id="job-aaaaaaaaaaaa", kind="contract_check")
        unit = UnitRecord(job.job_id, "contract-check", "contract_check", 0, {})
        out = run(self.processor(FakeRunner(good)).process(job, unit))
        report = json.loads(self.artifacts.get_bytes("jobs/job-aaaaaaaaaaaa/outputs/contract_report.json"))
        self.assertTrue(report["passed"])
        self.assertEqual([c["status"] for c in report["checks"]], ["pass", "pass", "pass"])
        self.assertAlmostEqual(out.cost_usd, 0.02)

        def leaky(task, n):                                         # Read executed outside the root, never denied
            if len(task.profile.builtin_tools) == 0:
                return good(task, n)
            return AgentRunStats(turns=1, tool_calls={"Read": 1})

        with self.assertRaises(AgentError) as cm:
            run(self.processor(FakeRunner(leaky)).process(job, unit))
        self.assertIn("deny_outside_root", str(cm.exception))


class FinalizerTests(Base):
    def prepare(self, statuses=None, logic_for=None):
        artifacts = InMemoryArtifactStore()
        plan = build_plan(JobOptions.from_dict({"include_tabular_sweep": False}), self.settings, self.corpus)
        job = JobRecord(job_id="job-aaaaaaaaaaaa", version_info={"service": "t"})
        runner = FakeRunner(lambda t, n: emit_default(t, logic=(logic_for or {}).get(t.unit_id)))
        proc = UnitProcessor(self.settings, self.corpus, runner, artifacts)
        units = []
        for spec in plan.units:
            unit = UnitRecord(job.job_id, spec.unit_id, spec.kind, spec.order, spec.spec)
            status = (statuses or {}).get(spec.unit_id, UnitStatus.SUCCEEDED.value)
            if status == UnitStatus.SUCCEEDED.value:
                out = run(proc.process(job, unit))
                unit.rules_key, unit.rules_count, unit.cost_usd, unit.turns = out.rules_key, out.rules_count, out.cost_usd, out.turns
            else:
                unit.error = "boom"
            unit.status = status
            units.append(unit)
        return job, units, artifacts

    def test_success_outputs(self):
        job, units, artifacts = self.prepare()
        result = Finalizer(self.settings, artifacts, self.corpus).finalize(job, units)
        self.assertEqual(result.status, "SUCCEEDED")
        self.assertEqual(result.summary["rules_raw"], 6)
        self.assertEqual(result.summary["rules_after_dedupe"], 6)
        self.assertEqual(result.summary["executable_rules"], 6)
        self.assertEqual(result.summary["automatable_rules"], 6)
        for key in ("catalogue", "executable_rules", "rules_raw", "rules_final", "mbs_facts", "units", "summary"):
            self.assertIn(key, result.artifacts, key)
        wb = load_workbook(BytesIO(artifacts.get_bytes(f"jobs/{job.job_id}/outputs/rule_catalogue.xlsx")))["rule_catalogue"]
        ids = [wb.cell(row=r, column=1).value for r in range(2, wb.max_row + 1)]
        self.assertEqual(ids, [f"RULE_{i:03d}" for i in range(1, 10)])          # 3 template rows + 6 new, sequential
        exec_ids = [r["rule_id"] for r in iter_jsonl(artifacts.get_bytes(f"jobs/{job.job_id}/outputs/executable_rules.jsonl"))]
        self.assertEqual(sorted(exec_ids), [f"RULE_{i:03d}" for i in range(4, 10)])   # join keys remapped to the workbook ids
        final = list(iter_jsonl(artifacts.get_bytes(f"jobs/{job.job_id}/outputs/rules_final.jsonl")))
        self.assertEqual([f["rule"]["rule_id"] for f in final], [f"RULE_{i:03d}" for i in range(4, 10)])
        self.assertEqual({f["rule"]["confidence"] for f in final}, {"needs_review"})
        summary = json.loads(artifacts.get_bytes(f"jobs/{job.job_id}/outputs/summary.json"))
        self.assertEqual(summary["status"], "SUCCEEDED")
        self.assertAlmostEqual(summary["cost_usd"], 0.3)

    def test_duplicates_across_units_are_collapsed(self):
        same = {f"acs-{c}": "identical logic" for c in ("0001", "0002")}
        job, units, artifacts = self.prepare(logic_for=same)
        result = Finalizer(self.settings, artifacts, self.corpus).finalize(job, units)
        self.assertEqual((result.summary["rules_raw"], result.summary["rules_after_dedupe"], result.summary["duplicates_dropped"]), (6, 5, 1))

    def test_partial_failed_and_cancelled_statuses(self):
        job, units, artifacts = self.prepare({"acs-0001": "FAILED"})
        fin = Finalizer(self.settings, artifacts, self.corpus)
        result = fin.finalize(job, units)
        self.assertEqual(result.status, "PARTIAL")
        self.assertEqual(result.summary["failures"], [{"unit_id": "acs-0001", "kind": "acs_topic", "error": "boom"}])
        job2, units2, artifacts2 = self.prepare({u: "FAILED" for u in ("acs-0001", "acs-0002", "acs-0010", "acs-0011", "acs-0110", "acs-0111")})
        result2 = Finalizer(self.settings, artifacts2, self.corpus).finalize(job2, units2)
        self.assertEqual(result2.status, "FAILED")
        self.assertNotIn("catalogue", result2.artifacts)
        job.cancel_requested = True
        self.assertEqual(fin.finalize(job, units).status, "CANCELLED")

    def test_missing_template_fails_closed_but_keeps_raw_ledger(self):
        settings = make_env(Path(self.tmp.name) / "nt", RULE_EXTRACTOR_TEMPLATE_PATH=str(Path(self.tmp.name) / "nope.xlsx"),
                            RULE_EXTRACTOR_SOURCE_DOCS_DIR=str(Path(self.tmp.name) / "nt" / "source_docs"))
        (Path(self.tmp.name) / "nt" / "template.xlsx").unlink()
        job, units, artifacts = self.prepare()
        result = Finalizer(settings, artifacts, self.corpus).finalize(job, units)
        self.assertEqual(result.status, "FAILED")
        self.assertIn("template not found", result.error)
        self.assertIn("rules_raw", result.artifacts)                  # nothing paid for is lost

    def test_base_executable_rules_are_carried_forward(self):
        base = Path(self.tmp.name) / "base.jsonl"
        base.write_text(json.dumps({"rule_id": "RULE_001", "flag_condition": {"has_diagnosis": "X00"},
                                    "correction_action": {"action": "flag_only"}, "automatable": False}) + "\n")
        settings = make_env(Path(self.tmp.name) / "bs", RULE_EXTRACTOR_BASE_RULES_PATH=str(base))
        job, units, artifacts = self.prepare()
        Finalizer(settings, artifacts, self.corpus).finalize(job, units)
        ids = [r["rule_id"] for r in iter_jsonl(artifacts.get_bytes(f"jobs/{job.job_id}/outputs/executable_rules.jsonl"))]
        self.assertEqual(ids[0], "RULE_001")
        self.assertEqual(len(ids), 7)

    def test_contract_check_job(self):
        artifacts = InMemoryArtifactStore()
        artifacts.put_bytes("jobs/job-aaaaaaaaaaaa/outputs/contract_report.json", b"{}")
        job = JobRecord(job_id="job-aaaaaaaaaaaa", kind="contract_check")
        unit = UnitRecord(job.job_id, "contract-check", "contract_check", 0, {}, status="SUCCEEDED")
        result = Finalizer(self.settings, artifacts, None).finalize(job, [unit])
        self.assertEqual(result.status, "SUCCEEDED")
        self.assertIn("contract_report", result.artifacts)


if __name__ == "__main__":
    unittest.main()
