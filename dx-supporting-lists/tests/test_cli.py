"""The command-line tools end to end, in --dry-run mode on generated sample data (no credentials, no network)."""
from __future__ import annotations

import contextlib
import csv
import io
import json
import os
import socket
import unittest
from pathlib import Path

from dx_support.cli import doctor
from dx_support.cli.main import main as cli_main
from dx_support.indexing.universe import load_universe
from dx_support.llm.bedrock import BedrockBackend
from dx_support.sample_data import write_sample_data
from dx_support.service.app import create_app
from dx_support.service.runtime import ServiceRuntime
from dx_support.settings import Settings
from dx_support.storage.checkpoint import Checkpoint
from dx_support.storage.locks import LOCK_KEY
from dx_support.storage.objectstore import LocalObjectStore, namespace
from tests.helpers import tmpdir
from tests.test_llm import ClientError, FakeConverse, tool_response

DATA = tmpdir() / "data"
write_sample_data(DATA)


def cli(*argv):
    """(exit code, stdout). SystemExit with a message counts as exit code 1 and the message is appended to stdout."""
    buf = io.StringIO()
    rc = 0
    with contextlib.redirect_stdout(buf):
        try:
            rc = cli_main(list(argv))
        except SystemExit as e:
            rc = e.code if isinstance(e.code, int) else 1
            if not isinstance(e.code, int) and e.code:
                buf.write(str(e.code))
    return rc, buf.getvalue()


class Fresh:
    """An isolated store and output folder per test."""

    def __init__(self):
        self.root = tmpdir()
        self.uri = (self.root / "store").as_uri()
        self.out = self.root / "compiled"
        self.common = ["--store-uri", self.uri, "--data-dir", str(DATA)]

    def build(self, *extra):
        return cli("build", "--dry-run", *self.common, "--yes", "--max-total-cost-usd", "5", "--category-concurrency", "4", *extra)

    def store(self, ns="runs/dry_run"):
        return namespace(LocalObjectStore(self.root / "store"), ns)


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


class PipelineCommandTests(unittest.TestCase):
    def setUp(self):
        self._saved = {k: os.environ.pop(k) for k in list(os.environ) if k.startswith("DX_")}

    def tearDown(self):
        for k in [k for k in os.environ if k.startswith("DX_")]:
            os.environ.pop(k)
        os.environ.update(self._saved)

    def test_build_status_compile_end_to_end(self):
        f = Fresh()
        rc, out = f.build()
        self.assertEqual(rc, 0, out)
        self.assertIn("'complete': 5", out)
        self.assertIn("'skipped_context_only': 1", out)
        rc, out = f.build()
        self.assertIn("Nothing to build", out)                                                       # a re-run buys nothing
        rc, out = cli("status", "--dry-run", *f.common)
        self.assertIn("Verdict    : FINISHED", out)
        rc, out = cli("compile", "--dry-run", *f.common, "--output-dir", str(f.out))
        self.assertEqual(rc, 0, out)
        files = {p.name for p in f.out.iterdir()}
        self.assertTrue({"supporting_lists.csv", "supporting_achi.csv", "supporting_mbs.csv", "supporting_prostheses.csv", "full_audit.csv", "excluded_codes.csv",
                         "supporting_lists_summary.xlsx", "dx_crosswalk_profiles.xlsx", "AI_GENERATED_NOTICE.txt"} <= files, files)
        self.assertNotIn("supporting_dx.csv", files)                                                 # the diagnosis family was not run
        lists = read_csv(f.out / "supporting_lists.csv")
        self.assertTrue(lists and all(r["support_tier"] in ("core", "broad") for r in lists))
        audit = read_csv(f.out / "full_audit.csv")
        self.assertTrue(any(r["support_tier"] == "not_supporting" for r in audit))                  # every verdict is in the audit
        self.assertTrue(all(r["merged"] in ("True", "False") for r in audit))
        notice = (f.out / "AI_GENERATED_NOTICE.txt").read_text()
        s = Settings.from_env({"DX_ENV": "test"})
        self.assertEqual(notice, f"{s.notice_line_1}\n{s.notice_line_2}\n")
        from openpyxl import load_workbook
        wb = load_workbook(f.out / "supporting_lists_summary.xlsx")
        self.assertEqual(wb.sheetnames, ["README", "by_category", "qa_flags", "ubiquity", "excluded_summary", "orphans", "run_info"])
        self.assertIn(s.notice_line_2, wb["by_category"].oddFooter.center.text)
        xw = load_workbook(f.out / "dx_crosswalk_profiles.xlsx")
        self.assertEqual(xw.sheetnames, ["achi_profiles", "mbs_profiles", "prosthesis_profiles"])
        header = [c.value for c in next(xw["achi_profiles"].iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, ["icd10_category", "description", "code", "code_description", "z_global", "p_global", "z_chapter", "p_chapter", "source"])
        rows = list(xw["achi_profiles"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(rows and all(isinstance(r[2], int) and r[8] == "llm_review" for r in rows))   # ACHI codes are numbers; statistics columns blank
        # scoped compile never overwrites a full one
        rc, out = cli("compile", "--dry-run", *f.common, "--output-dir", str(f.out), "--chapter", "7", "--merge-tiers", "core")
        self.assertEqual(rc, 0, out)
        self.assertTrue((f.out / "supporting_lists_chapter_VII.csv").exists())
        self.assertTrue(all(r["support_tier"] == "core" for r in read_csv(f.out / "supporting_lists_chapter_VII.csv")))
        rc, out = cli("compile", "--dry-run", *f.common, "--output-dir", str(f.out), "--chapter", "9-10", "--name-suffix", "pilot_a")
        self.assertTrue((f.out / "supporting_lists_pilot_a.csv").exists())
        rc, _ = cli("compile", "--dry-run", *f.common, "--name-suffix", "../escape")
        self.assertEqual(rc, 2)                                                                       # a suffix can never carry a path

    def test_spreadsheet_formulas_in_model_text_are_neutralised_in_every_output(self):
        f = Fresh()
        f.build()
        ck = Checkpoint(f.store())
        import asyncio
        asyncio.run(ck.write_verdicts([{"category": "I48", "family": "achi", "unit_id": "3880100", "support_tier": "core", "evidence_type": "treated",
                                        "rationale": "=HYPERLINK(\"http://evil\",\"x\")", "member_scope": "all", "source": "route", "model": "m", "ts": "2999-01-01T00:00:00"}]))
        ck.close()
        cli("compile", "--dry-run", *f.common, "--output-dir", str(f.out))
        row = next(r for r in read_csv(f.out / "supporting_achi.csv") if r["code"] == "3880100" and r["icd10_category"] == "I48")
        self.assertTrue(row["rationale"].startswith("'="), row["rationale"])
        self.assertTrue(next(r for r in read_csv(f.out / "full_audit.csv") if r["unit_id"] == "3880100" and r["icd10_category"] == "I48")["rationale"].startswith("'="))

    def test_a_changed_configuration_cannot_silently_continue_a_run(self):
        f = Fresh()
        self.assertEqual(f.build()[0], 0)
        rc, out = f.build("--thinking", "high")
        self.assertEqual(rc, 1)
        self.assertIn("different configuration", out)
        rc, out = f.build("--thinking", "high", "--accept-config-change")
        self.assertEqual(rc, 0, out)

    def test_a_second_run_against_the_same_namespace_is_refused_while_the_lock_is_held(self):
        f = Fresh()
        store = f.store()
        self.assertTrue(store.put_if_absent(LOCK_KEY, f"pid={os.getpid()} host={socket.gethostname()} started=2026-01-01T00:00:00".encode()))
        rc, out = f.build()
        self.assertEqual(rc, 1)
        self.assertIn("already exists", out)
        rc, out = cli("status", "--dry-run", *f.common)
        self.assertIn("RUNNING", out)                                                                 # this very process is alive
        store.delete(LOCK_KEY)
        store.put(LOCK_KEY, f"pid=999999 host={socket.gethostname()} started=2026-01-01T00:00:00".encode())
        rc, out = cli("status", "--dry-run", *f.common)
        self.assertIn("CRASHED OR KILLED", out)

    def test_for_service_prewarms_the_endpoint(self):
        f = Fresh()
        rc, out = f.build("--for-service", "--families", "achi,mbs")
        self.assertEqual(rc, 0, out)
        env = {"DX_ENV": "test", "DX_BACKEND": "fake", "DX_STORE_URI": f.uri, "DX_DATA_DIR": str(DATA)}
        rt = ServiceRuntime(Settings.from_env(env)).start()
        try:
            c = create_app(rt).test_client()
            r = c.post("/invocations", data=json.dumps({"categories": ["I48"], "families": ["achi", "mbs"]}), content_type="application/json")
            b = r.get_json()
            self.assertEqual((r.status_code, b["status"], b["usage"]["cache_hit"], b["usage"]["llm_calls"]), (200, "complete", True, 0))
            self.assertTrue(b["results"]["I48"]["families"]["achi"]["items"])
            r = c.post("/invocations", data=json.dumps({"categories": ["I48"], "families": ["prosthesis"]}), content_type="application/json")
            self.assertFalse(r.get_json()["usage"]["cache_hit"])                                       # a family that was not pre-built is computed, not faked
        finally:
            rt.shutdown(timeout=2)
        rc, out = f.build("--for-service", "--mode", "exhaustive")
        self.assertIn("only applies to the routed mode", out)

    def test_estimate_universe_and_chapters(self):
        f = Fresh()
        rc, out = cli("estimate", "--list-chapters", *f.common)
        self.assertIn("VII", out)
        self.assertIn("SKIP", out)
        rc, out = cli("estimate", "--show-universe", *f.common)
        self.assertIn("ACHI: 10 reviewable units", out)
        self.assertIn("placeholder_self_citing", out)
        rc, out = cli("estimate", *f.common, "--families", "all", "--chapter", "VII,IX")
        self.assertEqual(rc, 0, out)
        self.assertIn("[ROUTED", out)
        self.assertIn("[EXHAUSTIVE", out)
        rc, out = cli("estimate", *f.common, "--calibrate")
        self.assertIn("no usable pilot", out)
        f.build()
        rc, out = cli("estimate", "--dry-run", *f.common, "--calibrate")
        self.assertIn("Nothing left to build", out)
        rc, out = cli("estimate", "--dry-run", *f.common, "--families", "all")
        self.assertIn("dx for 5 categories", out)
        self.assertNotIn("CALIBRATED", out)                                                            # nobody has finished dx: there is no pilot to average

    def test_validation_pilots_run_end_to_end(self):
        f = Fresh()
        f.build()
        outdir = str(f.root / "validation")
        rc, out = cli("validate-retrieval", "--dry-run", *f.common, "--categories", "H25,I48", "--families", "achi,mbs", "--yes", "--output-dir", outdir)
        self.assertEqual(rc, 0, out)
        self.assertIn("RECALL of reference positives", out)
        self.assertIn("DRY RUN", out)
        rows = read_csv(Path(outdir) / "retrieval_validation_dryrun.csv")
        self.assertEqual({(r["category"], r["family"]) for r in rows}, {("H25", "achi"), ("H25", "mbs"), ("I48", "achi"), ("I48", "mbs")})
        rc, out = cli("validate-judge", "--dry-run", *f.common, "--n-pairs", "20", "--yes", "--output-dir", outdir)
        self.assertEqual(rc, 0, out)
        self.assertIn("REPEATABILITY", out)
        self.assertIn("kappa", out)
        gold = f.root / "gold.csv"
        gold.write_text("icd10_category,family,code,gold_tier\nI48,achi,3880100,core\nI48,achi,5500000,not_supporting\nH25,achi,4270100,broad\nZ99,achi,1,core\nI48,achi,0000000,core\nI48,achi,3880100,bogus\n")
        rc, out = cli("validate-judge", "--dry-run", *f.common, "--gold", str(gold), "--yes", "--output-dir", outdir)
        self.assertEqual(rc, 0, out)
        self.assertIn("GOLD mode", out)
        self.assertIn("unknown_category", out)
        self.assertIn("invalid_gold_tier", out)
        self.assertTrue((Path(outdir) / "judge_disagreements_dryrun.csv").exists())
        production = Checkpoint(f.store())
        self.assertFalse(any("validation" in k for k in LocalObjectStore(f.root / "store").list("runs/")))   # nothing leaked into a production namespace
        production.close()

    def test_validate_judge_rejects_a_malformed_gold_file(self):
        f = Fresh()
        f.build()
        bad = f.root / "bad.csv"
        bad.write_text("category,tier\nH25,core\n")
        rc, out = cli("validate-judge", "--dry-run", *f.common, "--gold", str(bad), "--yes")
        self.assertEqual(rc, 1)
        self.assertIn("must have columns", out)

    def test_dispatch_and_sample_data(self):
        self.assertEqual(cli()[0], 2)
        self.assertEqual(cli("help")[0], 0)
        rc, out = cli("nonsense")
        self.assertEqual(rc, 2)
        self.assertEqual(cli("version")[1].strip(), "2.0.0")
        d = tmpdir() / "s"
        self.assertEqual(cli("make-sample-data", "--out", str(d))[0], 0)
        rc, out = cli("make-sample-data", "--out", str(d))
        self.assertIn("already contains", out)
        self.assertEqual(cli("make-sample-data", "--out", str(d), "--force")[0], 0)
        self.assertEqual(len(load_universe(d).categories), 6)


class DoctorTests(unittest.TestCase):
    def settings(self, **kw):
        return Settings.from_env({"DX_ENV": "test", "DX_DATA_DIR": str(DATA), "DX_STORE_URI": tmpdir().as_uri()}).with_overrides(**kw)

    def test_offline_checks_pass_on_good_configuration(self):
        res = doctor.run_checks(self.settings(), offline=True)
        by = {n: (lvl, d) for n, lvl, d in res}
        self.assertEqual({lvl for lvl, _ in by.values()}, {"ok"}, by)
        self.assertIn("au.anthropic.claude-sonnet-5", by["models"][1])
        self.assertIn("5 categories".replace("5", "6"), by["reference data"][1])

    def test_missing_data_and_unknown_model_are_reported_not_raised(self):
        res = {n: lvl for n, lvl, _ in doctor.run_checks(self.settings(data_dir=Path("/nonexistent/x"), model="claude-sonet-5"), offline=True)}
        self.assertEqual((res["reference data"], res["models"]), ("fail", "fail"))

    def test_every_configured_model_is_probed_with_the_real_request_shape(self):
        client = FakeConverse([tool_response({"ok": True}), tool_response({"ok": True})])
        be = BedrockBackend(self.settings(route_model="claude-sonnet-5-5"), client=client)
        res = doctor.run_checks(be.settings, backend=be, skip_data=True)
        probes = [(n, lvl) for n, lvl, _ in res if n.startswith("bedrock probe")]
        self.assertEqual(sorted(n for n, _ in probes), ["bedrock probe au.anthropic.claude-sonnet-5", "bedrock probe au.anthropic.claude-sonnet-5-5"])
        self.assertEqual({lvl for _, lvl in probes}, {"ok"})
        self.assertEqual(len(client.calls), 2)
        self.assertTrue(all(c["toolConfig"]["toolChoice"] == {"auto": {}} for c in client.calls))

    def test_an_iam_or_profile_problem_is_a_failed_check_with_the_reason(self):
        client = FakeConverse([ClientError("AccessDeniedException", "User is not authorized to perform bedrock:InvokeModel", 403)])
        be = BedrockBackend(self.settings(), client=client)
        res = doctor.run_checks(be.settings, backend=be, skip_data=True)
        probe = next(r for r in res if r[0].startswith("bedrock probe"))
        self.assertEqual(probe[1], "fail")
        self.assertIn("fatal", probe[2])
        self.assertIn("not authorized", probe[2])
        wrong = FakeConverse([{"output": {"message": {"content": [{"text": "hello"}]}}, "stopReason": "end_turn", "usage": {}}])
        res = doctor.run_checks(BedrockBackend(self.settings(), client=wrong).settings, backend=BedrockBackend(self.settings(), client=wrong), skip_data=True)
        self.assertEqual(next(r for r in res if r[0].startswith("bedrock probe"))[1], "warn")


if __name__ == "__main__":
    unittest.main()
