"""dx-support compile -- turns whatever is currently stored into the output files. Safe to run at any point, including mid-run in another process: it
only ever READS the store and never writes to it or to the data directory.

Outputs (in --output-dir, default ./output/compiled/<namespace>):
  supporting_lists.csv          every merged (category, code) pair, all families, compact columns
  supporting_achi.csv           ACHI list (with rationale)
  supporting_mbs.csv            MBS list  (with rationale)
  supporting_prostheses.csv     prosthesis list -- one row per BILLING CODE, with its device type (the verdict is per device type)
  supporting_dx.csv             the diagnosis family's list with rationale (only when it has been run)
  full_audit.csv                EVERY verdict at unit level incl. not_supporting / unclear, with rationale -- for coder review
  supporting_lists_summary.xlsx README, by_category, qa_flags, ubiquity, excluded_summary, orphans, run_info
  dx_crosswalk_profiles.xlsx    THE RULE-BUILDING INPUT: one tab per code type (achi_profiles, mbs_profiles, prosthesis_profiles, and dx_profiles only if the
                                diagnosis family has been run) in the nine-column layout icd10_category, description, code, code_description, z_global, p_global,
                                z_chapter, p_chapter, source. ACHI and MBS codes are numbers; the four statistics columns are blank.
  orphan_codes.csv              reviewable units no category's list contains (only when coverage is high)
  excluded_codes.csv            every code that was never reviewed, with the reason
  AI_GENERATED_NOTICE.txt       the output notice (also in every workbook sheet's print footer); wording comes from DX_OUTPUT_NOTICE_LINE_1/2

Merge policy: --merge-tiers decides which tiers count as "supporting" (default core + broad). `--merge-tiers core` gives the strict list. Every verdict, merged
or not, stays in full_audit.csv.

Security: every free-text cell passes through safe_cell() so a description or rationale beginning with = + - @ cannot execute as a spreadsheet formula.
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Optional

from dx_support.cli.common import RUN_CONFIG_KEY, add_run_args, load_settings, namespace_path, open_run, with_chapter, with_name_suffix
from dx_support.constants import DEFAULT_MERGE_TIERS, FAMILIES, SUPPORT_TIERS
from dx_support.indexing.universe import load_universe
from dx_support.reference.chapters import CONTEXT_ONLY_CHAPTERS, parse_chapters
from dx_support.text import safe_cell

logger = logging.getLogger("dx_support.compile")

HEADER_FILL = "31006F"        # header colour of the workbooks (brand purple)
FOOTER_GRAY = "747679"
WORKBOOK_FONT = "Arial"
TIER_RANK = {t: i for i, t in enumerate(SUPPORT_TIERS)}
FILES = {"lists": "supporting_lists.csv", "achi": "supporting_achi.csv", "mbs": "supporting_mbs.csv", "prosthesis": "supporting_prostheses.csv",
         "dx": "supporting_dx.csv", "audit": "full_audit.csv", "summary": "supporting_lists_summary.xlsx", "crosswalk": "dx_crosswalk_profiles.xlsx",
         "orphans": "orphan_codes.csv", "excluded": "excluded_codes.csv", "notice": "AI_GENERATED_NOTICE.txt"}
CROSSWALK_COLUMNS = ["icd10_category", "description", "code", "code_description", "z_global", "p_global", "z_chapter", "p_chapter", "source"]
CROSSWALK_TIER_COLUMNS = ["support_tier", "evidence_type", "member_scope", "unit_id"]
CROSSWALK_SHEETS = {"achi": "achi_profiles", "mbs": "mbs_profiles", "prosthesis": "prosthesis_profiles", "dx": "dx_profiles"}
EXCEL_MAX_ROWS = 1_048_576


def _safe_suffix(value: str) -> str:
    try:
        with_name_suffix(Path("x.csv"), value)
    except ValueError as e:
        raise argparse.ArgumentTypeError(str(e))
    return value


def parse_args(argv=None):
    p = argparse.ArgumentParser(prog="dx-support compile", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    add_run_args(p, tag=True)
    p.add_argument("--output-dir", type=str, default=None, help="Where to write the files (default ./output/compiled/<namespace>).")
    p.add_argument("--chapter", "--chapters", type=parse_chapters, default=None, metavar="CHAPTERS",
                   help="Only these ICD-10-AM chapters. Writes chapter-suffixed filenames (..._chapter_IX, ..._chapters_IX_X) so a scoped compile never overwrites a full one.")
    p.add_argument("--merge-tiers", type=str, default=",".join(DEFAULT_MERGE_TIERS), help=f"Tiers that count as supporting (default '{','.join(DEFAULT_MERGE_TIERS)}'). Valid: {','.join(SUPPORT_TIERS)}.")
    p.add_argument("--prosthesis-extract", choices=["all", "latest"], default="all")
    p.add_argument("--include-coarse-prostheses", action="store_true", help="Only needed if the run used it and run_config.json is missing.")
    p.add_argument("--generic-ubiquity-threshold", type=int, default=30, help="A code supported for at least this many categories is flagged as possibly generic (default 30).")
    p.add_argument("--oversize-threshold", type=int, default=1500, help="Flag a category whose merged list in one family exceeds this many codes.")
    p.add_argument("--no-xlsx", action="store_true", help="Skip BOTH workbooks (the summary and the crosswalk profiles).")
    p.add_argument("--name-suffix", type=_safe_suffix, default=None, metavar="TEXT", help="Name every output file <n>_TEXT.<ext> instead of the automatic chapter suffix. Letters, digits, _ and - only (max 60).")
    p.add_argument("--no-crosswalk", action="store_true", help="Skip dx_crosswalk_profiles.xlsx.")
    p.add_argument("--source-label", type=str, default="llm_review", help="Value of the `source` column in the crosswalk workbook (default 'llm_review').")
    p.add_argument("--with-tier-columns", action="store_true", help="Append support_tier, evidence_type, member_scope and unit_id columns to each crosswalk tab.")
    p.add_argument("--on-screen-footer", action="store_true", help="Also write the output notice as visible rows under the crosswalk data (off by default: a program loading the sheet would read them as data).")
    args = p.parse_args(argv)
    tiers = tuple(t.strip() for t in args.merge_tiers.split(",") if t.strip())
    bad = set(tiers) - set(SUPPORT_TIERS)
    if bad or not tiers:
        p.error(f"--merge-tiers has unrecognised tier(s) {sorted(bad)}; valid: {list(SUPPORT_TIERS)}")
    args.merge_tiers = tiers
    return args


# ---------------------------------------------------------------------------
def _clean(value):
    """Free text -> inert text (formula-injection guard) and bounded for Excel's cell limit."""
    v = safe_cell(value)
    return v[:32000] if isinstance(v, str) else v


def _open_csv(path):
    fh = open(path, "w", newline="", encoding="utf-8-sig")   # BOM so Excel opens UTF-8 correctly
    return fh, csv.writer(fh)


def _style_sheet(ws, n_cols: int, n_rows: int, notice: tuple[str, str], widths: Optional[list[int]] = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=WORKBOOK_FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    body = Font(name=WORKBOOK_FONT, size=10)
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1, max_col=n_cols):
        for cell in row:
            cell.font = body
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    if n_rows:
        ws.auto_filter.ref = ws.dimensions
    ws.oddFooter.center.text = "\n".join(notice).replace("&", "&&")
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = FOOTER_GRAY
    r = n_rows + 4
    for i, line in enumerate(notice):
        cell = ws.cell(row=r + i, column=1, value=line)
        cell.font = Font(name=WORKBOOK_FONT, size=8, italic=True, color=FOOTER_GRAY)
        cell.alignment = Alignment(wrap_text=False, horizontal="left")


def _add_sheet(wb, title: str, header: list[str], rows: list[list], notice, widths: Optional[list[int]] = None) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    limit = 1_000_000
    for r in rows[:limit]:
        ws.append([_clean(v) for v in r])
    _style_sheet(ws, len(header), min(len(rows), limit), notice, widths)
    if len(rows) > limit:
        ws.cell(row=limit + 6, column=1, value=f"(truncated: {len(rows) - limit:,} further rows are in the CSV files)")


def _excel_code(family: str, code: str):
    """ACHI and MBS codes are NUMBERS in the earlier profile workbook, so they are written as numbers. A code with a leading zero, or any non-numeric
    code (every prosthesis code), stays text so nothing is lost."""
    if family in ("achi", "mbs") and code.isdigit() and not (len(code) > 1 and code.startswith("0")):
        return int(code)
    return code


def _code_sort_key(code: str):
    return (0, int(code), "") if code.isdigit() else (1, 0, code)


def write_crosswalk(path, rows_by_family: dict, source_label: str, with_tiers: bool, on_screen_footer: bool, notice: tuple[str, str]) -> dict:
    """The rule-building workbook: one tab per code type. Rows are sorted by category, then code. `description` is the category description (the
    lowest-sorted member code's description). Streaming write, so a full build does not need a workbook's worth of memory. The notice is in each sheet's
    print footer and the workbook properties; visible rows are opt-in because they would read back as data rows."""
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    wb.properties.title = "Supporting-list crosswalk profiles"
    wb.properties.description = " ".join(notice)
    header = CROSSWALK_COLUMNS + (CROSSWALK_TIER_COLUMNS if with_tiers else [])
    label = _clean(source_label)
    report: dict = {}
    for fam in [f for f in FAMILIES if f in rows_by_family]:          # the dx tab exists only when the diagnosis family was run
        ws = wb.create_sheet(CROSSWALK_SHEETS[fam])
        ws.oddFooter.center.text = "\n".join(notice).replace("&", "&&")
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.color = FOOTER_GRAY
        ws.append(header)
        rows = sorted(rows_by_family[fam], key=lambda r: (r[0], _code_sort_key(r[2])))
        written = 0
        for cat, cat_desc, code, code_desc, tier, evidence, scope, uid in rows:
            if written >= EXCEL_MAX_ROWS - 4:
                break
            row = [cat, _clean(cat_desc), _excel_code(fam, code), _clean(code_desc), None, None, None, None, label]
            if with_tiers:
                row += [tier, evidence, scope, uid]
            ws.append(row)
            written += 1
        truncated = len(rows) > written
        if truncated:
            logger.warning("%s was cut at Excel's row limit (%s of %s rows); the complete list is in the supporting_*.csv files.", CROSSWALK_SHEETS[fam], f"{written:,}", f"{len(rows):,}")
        if on_screen_footer:
            ws.append([None])
            for line in notice:
                ws.append([line])
        report[fam] = {"rows": written, "categories": len({r[0] for r in rows}), "truncated": truncated}
    wb.save(path)
    return report


def main(argv=None) -> int:
    args = parse_args(argv)
    settings = load_settings(args, quiet=True)
    ns = args.namespace
    ns_path = ns if ns.startswith(("runs/", "serve/")) else namespace_path(ns, args.dry_run, args.tag)
    run = open_run(settings, ns_path)
    ckpt = run.ckpt
    try:
        cfg = ckpt.read_json(RUN_CONFIG_KEY) or {}
        include_coarse = args.include_coarse_prostheses or bool(cfg.get("include_coarse_prostheses"))
        universe = load_universe(settings.data_dir, include_coarse, warm_indexes=False)
        categories, indices = universe.categories, universe.indices
        chapter = args.chapter
        chapter_label = ",".join(chapter) if chapter else None
        cats_in_scope = [c for c, i in categories.items() if chapter is None or i.chapter in chapter]
        scope_set = set(cats_in_scope)
        notice = (settings.notice_line_1, settings.notice_line_2)

        # ---- load verdicts (the newest record per key wins; objects are unordered, so recency comes from the record's own timestamp) ----
        verdicts: dict[tuple[str, str, str], tuple] = {}
        newest: dict[tuple[str, str, str], str] = {}
        n_unknown_unit = 0
        for r in ckpt.iter_verdicts(cats_in_scope if chapter else None):
            cat, fam, uid = r.get("category"), r.get("family"), r.get("unit_id")
            if cat not in scope_set:
                continue
            if fam not in indices or uid not in indices[fam].units or r.get("support_tier") not in SUPPORT_TIERS:
                n_unknown_unit += 1
                continue
            key, ts = (cat, fam, uid), str(r.get("ts", ""))
            if key in verdicts and ts < newest[key]:
                continue
            newest[key] = ts
            verdicts[key] = (r["support_tier"], r.get("evidence_type", ""), r.get("rationale", ""), r.get("member_scope", "all"), r.get("source", ""), r.get("model", ""))
        if not verdicts:
            sys.exit(f"No verdicts found in {run.store.uri} -- run `dx-support build` first{' (--dry-run)' if args.dry_run else ''}.")
        if n_unknown_unit:
            logger.warning("%d verdict(s) reference a unit that is not in the current index (e.g. coarse-label prostheses built with --include-coarse-prostheses -- pass it here too, or check run_config.json).", n_unknown_unit)
        logger.info("Loaded %s verdicts for %s categories in scope%s.", f"{len(verdicts):,}", f"{len(cats_in_scope):,}", f" (chapter {chapter_label})" if chapter else "")
        ckpt.load_status_all()

        dx_present = any(k[1] == "dx" for k in verdicts)
        ACTIVE = tuple(f for f in FAMILIES if f != "dx") + (("dx",) if dx_present else ())
        ORPHAN_FAMILIES = tuple(f for f in ACTIVE if f != "dx")

        out_dir = Path(args.output_dir) if args.output_dir else Path("output") / "compiled" / ns_path.replace("/", "_")
        out_dir.mkdir(parents=True, exist_ok=True)

        def out(name: str) -> Path:
            p = out_dir / FILES[name]
            return with_name_suffix(p, args.name_suffix) if args.name_suffix else with_chapter(p, chapter)

        tiers_merge = set(args.merge_tiers)
        f_audit, w_audit = _open_csv(out("audit"))
        f_all, w_all = _open_csv(out("lists"))
        f_achi, w_achi = _open_csv(out("achi"))
        f_mbs, w_mbs = _open_csv(out("mbs"))
        f_pros, w_pros = _open_csv(out("prosthesis"))
        f_dx = w_dx = None
        if dx_present:
            f_dx, w_dx = _open_csv(out("dx"))
            w_dx.writerow(["icd10_category", "category_title", "chapter", "supporting_category", "supporting_category_description", "support_tier", "relation_type", "member_scope", "source", "rationale"])
        w_audit.writerow(["icd10_category", "category_title", "chapter", "family", "unit_id", "unit_text", "n_codes", "support_tier", "evidence_type", "rationale", "member_scope", "source", "model", "merged"])
        common = ["icd10_category", "category_title", "chapter", "family", "code", "code_description", "support_tier", "evidence_type", "member_scope", "source", "unit_id"]
        w_all.writerow(common)
        w_achi.writerow(common[:-1] + ["rationale"])
        w_mbs.writerow(common[:-1] + ["rationale"])
        w_pros.writerow(["icd10_category", "category_title", "chapter", "code", "device_type", "support_tier", "evidence_type", "member_scope", "source", "unit_id", "suffixes",
                         "in_latest_extract", "label_resolution", "n_label_variants"])

        merged_counts: dict[tuple[str, str, str], int] = Counter()
        units_reviewed: dict[tuple[str, str], int] = Counter()
        scoped_verdicts: Counter = Counter()
        ubiquity: dict[str, Counter] = {f: Counter() for f in ACTIVE}
        merged_units: dict[str, set] = {f: set() for f in ACTIVE}
        tier_totals: Counter = Counter()
        models_seen: Counter = Counter()
        xwalk: dict[str, list] = {f: [] for f in ACTIVE}

        try:
            by_cat: dict[str, list] = defaultdict(list)
            for (cat, fam, uid), v in verdicts.items():
                by_cat[cat].append((fam, TIER_RANK[v[0]], uid, v))
            for cat in cats_in_scope:
                info = categories[cat]
                for fam, _rank, uid, v in sorted(by_cat.get(cat, []), key=lambda t: (ACTIVE.index(t[0]), t[1], t[2])):
                    tier, evidence, rationale, scope, source, model = v
                    index = indices[fam]
                    unit = index.units[uid]
                    merged = tier in tiers_merge
                    tier_totals[tier] += 1
                    models_seen[model] += 1
                    units_reviewed[(cat, fam)] += 1
                    scope_s = "all" if scope == "all" else ";".join(scope)
                    if scope != "all":
                        scoped_verdicts[cat] += 1
                    w_audit.writerow([_clean(x) for x in (cat, info.title, info.chapter, fam, uid, unit.text, len(unit.codes), tier, evidence, rationale, scope_s, source, model, merged)])
                    if not merged:
                        continue
                    codes = unit.codes
                    if fam == "prosthesis" and args.prosthesis_extract == "latest":
                        codes = tuple(c for c in codes if index.code_meta[c]["in_latest_extract"])
                    if not codes:
                        continue
                    merged_counts[(cat, fam, tier)] += len(codes)
                    ubiquity[fam][uid] += 1
                    merged_units[fam].add(uid)
                    for code in codes:
                        desc = index.code_text.get(code, "")
                        xwalk[fam].append((cat, info.members[0][1], code, desc, tier, evidence, scope_s, uid))
                        base = [cat, info.title, info.chapter, fam, code, desc, tier, evidence, scope_s, source]
                        w_all.writerow([_clean(x) for x in base + [uid]])
                        if fam == "prosthesis":
                            m = index.code_meta[code]
                            w_pros.writerow([_clean(x) for x in (cat, info.title, info.chapter, code, desc, tier, evidence, scope_s, source, uid, m["suffixes"], m["in_latest_extract"], m["label_resolution"], m["n_label_variants"])])
                        elif fam == "dx":
                            w_dx.writerow([_clean(x) for x in (cat, info.title, info.chapter, code, desc, tier, evidence, scope_s, source, rationale)])
                        else:
                            (w_achi if fam == "achi" else w_mbs).writerow([_clean(x) for x in base + [rationale]])
        finally:
            for fh in (f_audit, f_all, f_achi, f_mbs, f_pros, f_dx):
                if fh:
                    fh.close()

        # ---- excluded codes (never reviewed) ----
        f_ex, w_ex = _open_csv(out("excluded"))
        w_ex.writerow(["family", "code", "reason", "description"])
        excl_counts: Counter = Counter()
        for fam in ACTIVE:
            for code, reason in sorted(indices[fam].excluded.items()):
                w_ex.writerow([fam, code, reason, _clean(indices[fam].code_text.get(code, ""))])
                excl_counts[(fam, reason)] += 1
        f_ex.close()

        # ---- coverage, orphans ----
        eligible = [c for c in cats_in_scope if categories[c].chapter not in CONTEXT_ONLY_CHAPTERS]
        status_of = {c: ckpt.status.get(c, {}).get("status", "not_run") for c in cats_in_scope}
        n_complete = sum(1 for c in eligible if status_of[c] == "complete")
        coverage = n_complete / len(eligible) if eligible else 0.0
        dx_missing = [c for c in eligible if "dx" not in (ckpt.status.get(c, {}).get("families_complete") or [])] if dx_present else []
        orphan_rows: list[tuple] = []
        orphan_note = ""
        if coverage >= 0.9:
            for fam in ORPHAN_FAMILIES:
                for uid, unit in sorted(indices[fam].units.items()):
                    if uid not in merged_units[fam]:
                        orphan_rows.append((fam, uid, unit.text, len(unit.codes)))
            f_or, w_or = _open_csv(out("orphans"))
            w_or.writerow(["family", "unit_id", "unit_text", "n_codes"])
            for r in orphan_rows:
                w_or.writerow([_clean(x) for x in r])
            f_or.close()
        else:
            orphan_note = f"Orphan report skipped: only {coverage:.0%} of eligible categories are complete (needs >= 90%) -- with partial coverage almost every code looks like an orphan."

        # ---- by-category summary + QA flags ----
        flags: list[list] = []
        by_cat_rows: list[list] = []
        n_not_run = sum(1 for c in eligible if status_of[c] == "not_run")
        if n_not_run > 50:
            flags.append(["(many)", "not_run", f"{n_not_run:,} eligible categories have not been processed yet (see by_category status)"])
        if dx_missing:
            flags.append(["(dx)", "dx_incomplete", f"{len(dx_missing):,} eligible categories have no completed diagnosis-family pass (re-run `dx-support build --families dx`); their dx_profiles rows are missing"])
        for cat in cats_in_scope:
            info = categories[cat]
            st = ckpt.status.get(cat, {})
            status = status_of[cat]
            cell = {(f, t): merged_counts.get((cat, f, t), 0) for f in ACTIVE for t in ("core", "broad")}
            row_flags = []
            if status in ("partial", "not_run") and info.chapter not in CONTEXT_ONLY_CHAPTERS:
                row_flags.append("not_complete")
                if status == "partial" or n_not_run <= 50:
                    flags.append([cat, "not_complete", f"status={status}" + (f"; {st.get('error')}" if st.get("error") else "")])
            for fam, fs in (st.get("families", {}) or {}).items():
                if fs.get("truncated_nodes") or fs.get("family_truncated"):
                    row_flags.append(f"truncated_{fam}")
                    flags.append([cat, f"truncated_{fam}", f"truncated_nodes={fs.get('truncated_nodes')}, family_truncated={fs.get('family_truncated')} -- some candidates were cut by the BM25 cap; consider --mode exhaustive for this category"])
            if status == "complete":
                totals = {f: cell[(f, "core")] + cell[(f, "broad")] for f in ACTIVE}
                if sum(totals.values()) == 0:
                    row_flags.append("empty_all_families")
                    flags.append([cat, "empty_all_families", "complete but no supporting code in any family (legitimate for some categories -- review)"])
                for f, n in totals.items():
                    if n > args.oversize_threshold:
                        row_flags.append(f"oversized_{f}")
                        flags.append([cat, f"oversized_{f}", f"{n:,} supporting codes in {f} (> {args.oversize_threshold:,}); typical for generic orthopaedic hardware"])
            if scoped_verdicts.get(cat):
                row_flags.append("member_scoped")
            by_cat_rows.append([cat, info.chapter, info.title, len(info.members), info.is_large, status, *[cell[(f, t)] for f in ACTIVE for t in ("core", "broad")],
                                *[units_reviewed.get((cat, f), 0) for f in ACTIVE], "; ".join(row_flags)])

        ubi_rows: list[list] = []
        for fam in ACTIVE:
            for uid, n in ubiquity[fam].most_common(150):
                ubi_rows.append([fam, uid, indices[fam].units[uid].text, n, n >= args.generic_ubiquity_threshold])
        n_generic = sum(1 for r in ubi_rows if r[4])

        cost_by_stage: Counter = Counter()
        n_attempts: Counter = Counter()
        for r in ckpt.iter_run_log():
            cost_by_stage[r.get("stage", "?")] += r.get("cost_usd") or 0.0
            n_attempts[r.get("stage", "?")] += 1
        status_counts = Counter(status_of.values())
        info_rows = [
            ["store", run.store.uri], ["dry run (offline stub -- NOT clinical)", bool(args.dry_run)], ["mode", cfg.get("mode", "routed")], ["fingerprint", cfg.get("fingerprint", "")],
            ["chapter scope", chapter_label or "all"], ["merge policy (tiers counted as supporting)", ", ".join(args.merge_tiers)], ["prosthesis extract policy", args.prosthesis_extract],
            ["categories in scope", len(cats_in_scope)], ["category status", dict(status_counts)], ["eligible categories complete", f"{n_complete} of {len(eligible)} ({coverage:.0%})"],
            ["verdicts loaded", len(verdicts)], ["verdicts by tier", dict(tier_totals)], ["models seen in verdicts", dict(models_seen)],
            ["total cost over all logged attempts (USD)", round(sum(cost_by_stage.values()), 2)], ["cost by stage (USD)", {k: round(v, 2) for k, v in cost_by_stage.items()}],
            ["attempts by stage", dict(n_attempts)], ["codes flagged possibly generic (top-150 view)", n_generic], ["orphans", orphan_note or f"{len(orphan_rows):,} units (orphan_codes.csv)"],
            ["note", "A verdict is a screening signal, not a coding determination. A qualified clinical coder should review the lists."],
        ]
        if args.dry_run:
            info_rows.insert(0, ["WARNING", "DRY RUN: verdicts come from an offline stub that matches words only. They carry no clinical meaning."])

        n_incomplete = sum(1 for c in eligible if status_of[c] != "complete")
        xwalk_report = None
        if not args.no_xlsx and not args.no_crosswalk:
            xwalk_report = write_crosswalk(out("crosswalk"), xwalk, args.source_label, args.with_tier_columns, args.on_screen_footer, notice)

        if not args.no_xlsx:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            readme = [
                ["Purpose", "Supporting lists: ACHI codes, MBS items and prosthesis billing codes whose presence on a hospital episode indicates that the ICD-10-AM category was treated, investigated or examined."],
                ["Tiers", "core = specific to the category's condition/site; broad = non-specific treatment/investigation/examination consistent with it; unclear; not_supporting."],
                ["Merged", f"Supporting lists contain: {', '.join(args.merge_tiers)}. Every verdict is in full_audit.csv."],
                ["by_category", "One row per ICD-10-AM category: status, supporting-code counts per family and tier, units reviewed, flags."],
                ["qa_flags", "Categories needing attention: incomplete, truncated candidate sets, empty or oversized lists."],
                ["ubiquity", "Codes that support many categories -- candidates for being generic (data-driven)."],
                ["excluded_summary", "Codes never reviewed, by reason (full list: excluded_codes.csv)."],
                ["Prostheses", "Verdicts are per device type; each applies to every billing code in the group (the file has no product names)."],
                ["Review", "A qualified clinical coder must review these lists before any production use."],
            ]
            if dx_present:
                readme.insert(-1, ["Diagnosis family", "dx rows: for a TARGET category coded as an additional diagnosis, the OTHER categories whose presence in the same episode makes it clinically plausible "
                                                       "(relation_type: underlying_condition, complication_or_manifestation, associated_comorbidity). A plausibility screen, not the coding standard's test."])
            _add_sheet(wb, "README", ["Topic", "Description"], readme, notice, [24, 140])
            _add_sheet(wb, "by_category", ["category", "chapter", "title", "n_member_codes", "large_category", "status"] + [f"{f}_{t}" for f in ACTIVE for t in ("core", "broad")]
                       + [f"{f}_units_reviewed" for f in ACTIVE] + ["flags"], by_cat_rows, notice, [10, 9, 60, 12, 12, 18] + [13] * (2 * len(ACTIVE)) + [20] * len(ACTIVE) + [40])
            _add_sheet(wb, "qa_flags", ["category", "flag", "detail"], flags or [["", "none", "no flags raised"]], notice, [10, 24, 120])
            _add_sheet(wb, "ubiquity", ["family", "unit_id", "text", "n_categories_supported", f"possibly_generic(>={args.generic_ubiquity_threshold})"], ubi_rows, notice, [12, 18, 100, 20, 22])
            _add_sheet(wb, "excluded_summary", ["family", "reason", "n_codes"], [[f, r, n] for (f, r), n in sorted(excl_counts.items())], notice, [14, 40, 10])
            _add_sheet(wb, "orphans", ["family", "n_orphan_units"] if orphan_rows else ["note"],
                       [[f, sum(1 for r in orphan_rows if r[0] == f)] for f in ORPHAN_FAMILIES] if orphan_rows else [[orphan_note or "no orphans"]], notice, [14, 20] if orphan_rows else [120])
            _add_sheet(wb, "run_info", ["item", "value"], [[k, json.dumps(v) if isinstance(v, (dict, list)) else v] for k, v in info_rows], notice, [48, 100])
            wb.save(out("summary"))

        out("notice").write_text("\n".join(notice) + "\n", encoding="utf-8")

        print("=" * 72)
        print(f"Categories in scope: {len(cats_in_scope):,} | status: {dict(status_counts)}")
        print(f"Merge policy: {', '.join(args.merge_tiers)} | verdicts: {len(verdicts):,} {dict(tier_totals)}")
        for fam in ACTIVE:
            n = sum(v for (c, f, t), v in merged_counts.items() if f == fam)
            print(f"  {fam:>10}: {n:,} (category, code) pairs in the supporting list; {len(merged_units[fam]):,} distinct units used")
        print(f"QA flags: {len(flags)} | possibly generic codes (top-150): {n_generic} | {orphan_note or f'orphan units: {len(orphan_rows):,}'}")
        print(f"Written to: {out_dir}{f' (chapter {chapter_label})' if chapter else ''}")
        for key in ("lists", "achi", "mbs", "prosthesis") + (("dx",) if dx_present else ()) + ("audit", "excluded") + (() if args.no_xlsx else ("summary",) + (() if args.no_crosswalk else ("crosswalk",))):
            print(f"  - {out(key).name}")
        if xwalk_report:
            print(f"Rule-building workbook {out('crosswalk').name}: " + " | ".join(f"{CROSSWALK_SHEETS[f]} {xwalk_report[f]['rows']:,} rows, {xwalk_report[f]['categories']} categories" for f in ACTIVE))
            if n_incomplete:
                print(f"  WARNING: {n_incomplete} eligible categor{'y is' if n_incomplete == 1 else 'ies are'} not complete: re-run `dx-support build` before building rules from this workbook.")
            if dx_missing:
                print(f"  WARNING: the diagnosis family is not complete for {len(dx_missing)} eligible categories: re-run `dx-support build --families dx` before building rules from the dx_profiles tab.")
        print("The reference data files were not modified. Review full_audit.csv before relying on any list.")
        print("=" * 72)
        return 0
    finally:
        run.close()


if __name__ == "__main__":
    raise SystemExit(main())
